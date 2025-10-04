import os
import json
import pandas as pd
import numpy as np
from datetime import datetime

# Try to import python-pptx (pip install python-pptx)
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.dml.color import RGBColor
except Exception as e:
    raise SystemExit(
        "python-pptx is required for this script.\n"
        "Install it with: pip install python-pptx\n"
        f"Import error: {e}"
    )

RESULTS_DIR = "outputs/report"
SUIT_DIR = "outputs/suitability_ml_lgbm"
CAL_CA_JSON = "outputs/calibration/bayes_joint_best.json"
VAL_CA_JSON = "outputs/validation/validation_results.json"
CAL_ABM_JSON = "outputs/calibration_abm/abm_best.json"
VAL_ABM_JSON = "outputs/validation/validation_abm_results.json"
COSTS_CSV = "outputs/costs/costs_summary_2050.csv"
UNC_CSV = "outputs/uncertainty/uncertainty_summary.csv"

# Figures created by reporting.py (recommended to run that first)
SUIT_FIG = os.path.join(RESULTS_DIR, "suitability", "feature_importance.png")
CA_CM_FIG = os.path.join(RESULTS_DIR, "ca", "confusion_matrix.png")
ABM_COMP_FIG = os.path.join(RESULTS_DIR, "abm", "abm_vs_ca.png")
SC_COSTS_FIG = os.path.join(RESULTS_DIR, "scenarios", "cost_breakdown.png")
SC_UNC_FIG = os.path.join(RESULTS_DIR, "scenarios", "uncertainty.png")

DECK_OUT = os.path.join(RESULTS_DIR, "Urban_Modeling_Report.pptx")


# ---------------------------
# Helpers
# ---------------------------
def add_title_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle:
        slide.placeholders[1].text = subtitle


def add_title_only(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = title
    return slide


def add_bullets(slide, items):
    left = Inches(0.6); top = Inches(1.6); width = Inches(11); height = Inches(5)
    tf = slide.shapes.add_textbox(left, top, width, height).text_frame
    for i, it in enumerate(items):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = str(it)
        p.level = 0


def add_image(slide, image_path, max_height=5.0, max_width=10.5, left_in=0.5, top_in=1.5):
    if not os.path.exists(image_path):
        add_bullets(slide, [f"(missing) {os.path.basename(image_path)}"])
        return
    pic = slide.shapes.add_picture(image_path, Inches(left_in), Inches(top_in))
    # scale down if needed, keep aspect
    if pic.height > Inches(max_height):
        ratio = Inches(max_height) / pic.height
        pic.height = Inches(max_height)
        pic.width = int(pic.width * ratio)
    if pic.width > Inches(max_width):
        ratio = Inches(max_width) / pic.width
        pic.width = Inches(max_width)
        pic.height = int(pic.height * ratio)


def fmt_val(val):
    """Pretty formatting for table cells."""
    # Convert sequences/arrays to short comma-separated string
    if isinstance(val, (list, tuple, pd.Series, pd.Index, np.ndarray)):
        return ", ".join(map(str, val))
    # Empty / NaN
    try:
        if pd.isna(val):
            return ""
    except Exception:
        pass
    # Numbers
    if isinstance(val, (float, np.floating)):
        return f"{val:.3f}" if abs(val) < 1000 else f"{val:,.0f}"
    if isinstance(val, (int, np.integer)):
        return f"{val:,}"
    return str(val)


def add_table(slide, df: pd.DataFrame, title=None):
    if df is None or df.empty:
        add_bullets(slide, ["(no data)"])
        return
    if title:
        add_bullets(slide, [title])

    rows, cols = df.shape
    left = Inches(0.5); top = Inches(1.5); width = Inches(10.0); height = Inches(0.8 + 0.3 * (rows + 1))
    table = slide.shapes.add_table(rows + 1, cols, left, top, width, height).table
    # header
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)
    # rows
    for i in range(rows):
        for j in range(cols):
            table.cell(i + 1, j).text = fmt_val(df.iloc[i, j])


def safe_load_json(path):
    if not os.path.exists(path):
        return None
    with open(path, "r") as f:
        return json.load(f)


def safe_load_csv(path):
    return pd.read_csv(path) if os.path.exists(path) else None


# ---------------------------
# Build Deck
# ---------------------------
def run():
    os.makedirs(RESULTS_DIR, exist_ok=True)
    prs = Presentation()

    # 1) Title
    add_title_slide(
        prs,
        "Spatially Modelling – Cost of Sprawl",
        f"Urban growth suitability, CA/ABM simulation, scenarios & costs\nGenerated: {datetime.now():%Y-%m-%d %H:%M}"
    )

    # 2) Overview slide
    s = add_title_only(prs, "Project Overview")
    bullets = [
        "Data-driven suitability (LGBM) using NDVI, LST, population, LULC, distances, slope.",
        "Cellular Automata (CA) with anisotropic + multi-scale neighborhoods.",
        "ABM layer atop CA: developer profitability + household utility.",
        "Calibration & validation on 2000→2010 and 2010→2020 (FoM, Quantity/Allocation).",
        "Policy scenarios to 2050 + cost & uncertainty analyses."
    ]
    add_bullets(s, bullets)

    # 3) Suitability
    s = add_title_only(prs, "Suitability (LGBM)")
    add_image(s, SUIT_FIG)
    metrics = safe_load_json(os.path.join(RESULTS_DIR, "suitability", "metrics_summary.json"))
    if metrics:
        s2 = add_title_only(prs, "Suitability – Validation Metrics")
        dfm = pd.DataFrame([metrics])
        add_table(s2, dfm)

    # 4) CA calibration & validation
    ca_best = safe_load_json(CAL_CA_JSON)
    if ca_best:
        s = add_title_only(prs, "CA Calibration – Best Parameters")
        df_params = pd.DataFrame([ca_best.get("params", {})])
        add_table(s, df_params)

        if "metrics" in ca_best:
            s = add_title_only(prs, "CA Calibration – Metrics")
            dfm = pd.DataFrame([ca_best["metrics"]])
            add_table(s, dfm)

    ca_val = safe_load_json(VAL_CA_JSON)
    if ca_val:
        s = add_title_only(prs, "CA Validation (2010→2020)")
        cols = ["FoM","precision","recall","f1","true_positive","false_positive","false_negative","true_negative","sim_new","obs_new"]
        dfv = pd.DataFrame([{k: ca_val.get(k, None) for k in cols}])
        add_table(s, dfv)
        s2 = add_title_only(prs, "CA – Confusion Matrix")
        add_image(s2, CA_CM_FIG)

    # 5) ABM calibration & validation
    abm_best = safe_load_json(CAL_ABM_JSON)
    if abm_best:
        s = add_title_only(prs, "ABM Calibration – Best Weights")
        dfp = pd.DataFrame([abm_best.get("params", {})])
        add_table(s, dfp)

    abm_val = safe_load_json(VAL_ABM_JSON)
    if abm_val:
        s = add_title_only(prs, "ABM vs CA-only (Validation 2010→2020)")
        add_image(s, ABM_COMP_FIG)
        s2 = add_title_only(prs, "ABM Validation – Metrics")
        cols = ["FoM","precision","recall","f1","sim_new","obs_new"]
        df_abm = pd.DataFrame([{k: abm_val.get(k, None) for k in cols}])
        add_table(s2, df_abm)

    # 6) Scenarios – Costs & Uncertainty
    costs = safe_load_csv(COSTS_CSV)
    unc = safe_load_csv(UNC_CSV)

    if costs is not None:
        s = add_title_only(prs, "Scenario Costs (2020→2050)")
        add_image(s, SC_COSTS_FIG)
        s2 = add_title_only(prs, "Scenario Cost Summary (Million Rs.)")
        keep_cols = [
            "scenario","infra_cost_million","service_cost_million","env_cost_million",
            "revenue_million","total_cost_million","fia_million","new_cells","pop_new"
        ]
        dfc = costs[keep_cols] if set(keep_cols).issubset(costs.columns) else costs
        add_table(s2, dfc)

    if unc is not None:
        s = add_title_only(prs, "Scenario Uncertainty (Urban Cells by 2050)")
        add_image(s, SC_UNC_FIG)
        s2 = add_title_only(prs, "Uncertainty Summary")
        add_table(s2, unc)

    # 7) Notes & Next Steps
    s = add_title_only(prs, "Notes & Next Steps")
    next_steps = [
        "Refine suitability features (e.g., add amenities/heat risk, better land price).",
        "Per-year growth schedules matching observed/forecast demand.",
        "Region-wise growth caps + corridor policies as constraints.",
        "Tighter CA/ABM calibration with spatial cross-validation.",
        "Stakeholder-weighted multi-criteria decision analysis for scenario ranking."
    ]
    add_bullets(s, next_steps)

    prs.save(DECK_OUT)
    print(f"\n✅ Saved PowerPoint: {DECK_OUT}")


if __name__ == "__main__":
    run()
